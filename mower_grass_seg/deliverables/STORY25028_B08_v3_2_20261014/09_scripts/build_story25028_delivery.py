"""Build an evidence-backed DOCX/HTML delivery package without changing models."""
import csv
import argparse
import hashlib
import html
import json
from pathlib import Path
import shutil
import zipfile

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/root/cpp_project/am970/model')
PROJECT = Path('/mnt/e/cpp_project/lawn_mowing_robot_yolov8/mower_grass_seg')
RESULT = ROOT/'results/b08_v3_2_consistency_20261014'
OUT = Path('/root/cpp_project/am970/deliverables/STORY25028_B08_v3_2_20261014')


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def copy(source, relative):
    target = OUT/relative
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, target)
    assert sha(source) == sha(target)
    return relative


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--refresh',action='store_true')
    args=parser.parse_args()
    OUT.mkdir(parents=True, exist_ok=args.refresh)
    summary = json.loads((RESULT/'summary.json').read_text())
    accuracy = json.loads((RESULT/'accuracy_and_timing.json').read_text())
    native = ROOT/'results/b08_v3_2_input_shapes_20261014'
    native_summary = json.loads((native/'summary.json').read_text())
    models = RESULT/'deliverables/models'
    for p in models.iterdir():
        if p.is_file():
            copy(p, '02_models/'+p.name)
    expected = summary['model_sha256']
    for name, filename in [('pytorch','b08_v3_2.pt'),('onnx','b08_v3_2.onnx'),('npu','network_binary.nb')]:
        assert sha(OUT/'02_models'/filename) == expected[name]
    training = PROJECT/'runs/segment/B08_v3_2'
    for name in ['args.yaml','results.csv','results.png','MaskPR_curve.png','confusion_matrix.png']:
        copy(training/name, '03_training/'+name)
    copy(PROJECT/'TRAINING_LOG.md', '03_training/TRAINING_LOG_snapshot.md')
    work = ROOT/'work/b08_v3_2_20261014'
    for name in ['import.log','quantize_initial.log','quantize_hybrid.log','export.log','provenance.json']:
        copy(work/name, '04_conversion/'+name)
    for name in ['summary.json','accuracy_and_timing.json','per_image.csv','versions.json',
                 'board_run.log','board_checksums.txt','deployment_before.txt','deployment_after.txt']:
        copy(RESULT/name, '05_evaluation/current/'+name)
    for folder in ['comparisons','masks','vs_annotation','reference','vs_b08_v3_npu']:
        for p in (RESULT/'deliverables'/folder).rglob('*'):
            if p.is_file():
                copy(p, str(Path('05_evaluation/current')/folder/p.relative_to(RESULT/'deliverables'/folder)))
    for name in ['summary.json','per_image.csv']:
        copy(native/name,'05_evaluation/native_pt/'+name)
    for run in ['coord_chain_fp16_20261013','coord_fp16_20261013','crop_fix_20261013']:
        copy(ROOT/'results'/run/'summary.json','05_evaluation/history/'+run+'.json')
    copy(ROOT/'results/coord_chain_fp16_20261013/README.md','05_evaluation/history/coordinate_precision_notes.md')
    for run in ['b08_v3_consistency_20261014','pt_input_shapes_original_source_20261013','b08_v3_input_shapes_20261014']:
        copy(ROOT/'results'/run/'summary.json','05_evaluation/history/'+run+'.json')
    copy(ROOT/'results/b08_v3_consistency_20261014/accuracy_and_timing.json',
         '05_evaluation/history/b08_v3_accuracy_and_timing.json')

    figures = [
        ('009_coordinate_chain', ROOT/'results/coord_chain_fp16_20261013/009.jpg'),
        ('018_training_before_after', native/'vs_previous/018.jpg'),
        ('018_final_backends', RESULT/'comparisons/018.jpg'),
        ('042_npu_before_after', RESULT/'deliverables/vs_b08_v3_npu/042.jpg'),
        ('042_final_backends', RESULT/'comparisons/042.jpg'),
        ('044_training_before_after', native/'vs_previous/044.jpg'),
        ('044_npu_before_after', RESULT/'deliverables/vs_b08_v3_npu/044.jpg'),
        ('022_remaining_difference', RESULT/'comparisons/022.jpg'),
        ('023_remaining_difference', RESULT/'comparisons/023.jpg'),
    ]
    for name,p in figures:
        copy(p,'06_cases/'+name+'.jpg')
    for sid in ['009','018','042','044','022','023']:
        copy(RESULT/'reference'/f'{sid}_detections.json','06_cases/'+sid+'_detections.json')
    sample_map = {s['id']:s for s in summary['samples']}
    (OUT/'06_cases/sample_index.json').write_text(json.dumps(
        {k:sample_map[k] for k in ['009','018','042','044','022','023']},indent=2,ensure_ascii=False))

    data = OUT/'10_data_snapshot'
    data.mkdir(exist_ok=args.refresh)
    snapshots = []
    for s in native_summary['samples']:
        sid=s['id']; mapped=sample_map[sid]
        source=PROJECT/'data/test_train25_val25_b07_v4'/mapped['file']
        assert sha(source)==s['image_sha256']
        copy(source,'10_data_snapshot/internal50/images/'+mapped['file'])
        copy(native/'masks'/f'{sid}_gt.png','10_data_snapshot/internal50/gt_masks/'+sid+'.png')
        label=PROJECT/'data/yolo_seg/labels'/s['split']/Path(s['file']).with_suffix('.txt')
        label_matches=sha(label)==s['label_sha256']
        if label_matches:
            copy(label,'10_data_snapshot/internal50/labels/'+Path(mapped['file']).with_suffix('.txt').name)
        snapshots.append(dict(**s,label_snapshot_matches=label_matches))
    (data/'internal50/provenance.json').write_text(json.dumps(snapshots,indent=2))
    copy(PROJECT/'data/test_train25_val25_b07_v4/manifest.csv','10_data_snapshot/internal50/manifest.csv')
    provenance=json.loads((work/'provenance.json').read_text())['sha256']
    calibration=[]
    for line in (ROOT/'calib_mower/dataset.txt').read_text().splitlines():
        p=Path(line)
        assert sha(p)==provenance[str(p)]
        rel=copy(p,'10_data_snapshot/calibration100/images/'+p.name)
        calibration.append(dict(original_path=str(p),package_path=rel,sha256=sha(p)))
    (data/'calibration100/manifest.json').write_text(json.dumps(calibration,indent=2))
    copy(ROOT/'calib_mower/dataset.txt','10_data_snapshot/calibration100/original_dataset.txt')
    for name in ['compare_mower_backends.py','evaluate_mower_pt_input_shapes.py','convert_b08_v3.py',
                 'convert_mower_coord_fp16.py','export_mower_onnx.py','report_b08_v3.py',
                 'test_compare_mower_backends.py','build_story25028_delivery.py']:
        copy(ROOT/'scripts'/name,'09_scripts/'+name)
    for name in ['install_models.sh','run_board.sh']:
        copy(work/name,'08_deployment/'+name)
    (OUT/'08_deployment/installation_receipt.json').write_text(json.dumps(dict(
        evidence='Recorded successful installation and checksum verification on 2026-10-14; not a new live board check at report generation.',
        model='B08_v3_2',sha256=expected['npu'],
        targets=['/root/deploy_mower_cam_vknn/model/network_binary.nb','/root/mower_uint8_trial/network_binary.nb'],
        backup='/root/model_backups/before_b08_v3_2_20261014_8421',
        scope='Only model files replaced; camera executable and launch parameters unchanged.',
        note='05_evaluation deployment_before/after files describe the preceding isolated test, not this later installation.'),indent=2))
    cam = PROJECT/'data/cam_detect/b08_v3_2_camera_20261014_run8512'
    for name in ['run.log','timing.csv','camera_debug.txt','frame_19700101_120239_cam0.jpg']:
        copy(cam/name,'08_deployment/indoor_camera_smoke/'+name)
    field=OUT/'07_field_test_PENDING'
    field.mkdir(exist_ok=args.refresh)
    with (field/'field_records_template.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f)
        w.writerow(['序号','日期','地点/序列','光照','场景','原图文件','标注文件','模型SHA256','帧数',
                    '运行秒数','平均FPS','误检记录','漏检记录','异常/重启','电源','结论','复核人'])
        w.writerow(['待补充']+['']*16)
    (field/'STATUS.txt').write_text('状态：未完成，待补充独立外场数据测试。\n'
        '请勿将内部50图或室内无草坪运行结果填作外场验收。\n'
        '补充原图、人工标注、结果图、run.log、timing.csv、场景清单和签字结论。\n',encoding='utf-8')

    doc=Document()
    sec=doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.6);sec.bottom_margin=Cm(1.6);sec.left_margin=Cm(1.6);sec.right_margin=Cm(1.6)
    for sty in ['Normal','Title','Heading 1','Heading 2','Caption']:
        style=doc.styles[sty]
        style.font.name='Microsoft YaHei'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
    doc.styles['Normal'].font.size=Pt(10)
    doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    doc.styles['Normal'].paragraph_format.line_spacing=1.15
    doc.styles['Heading 1'].font.size=Pt(17)
    doc.styles['Heading 2'].font.size=Pt(12)
    sec.header.paragraphs[0].text='STORY #25028 | B08_v3_2 | 阶段交付 · 2026-10-14'
    footer=sec.footer.paragraphs[0];footer.alignment=2
    footer.add_run('第 ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    footer.add_run(' 页')
    blocks=[]
    def heading(text,level=1,page=False):
        if page:doc.add_page_break()
        doc.add_heading(text,level)
        blocks.append(f'<h{level}>{html.escape(text)}</h{level}>')
    def p(text):
        doc.add_paragraph(text)
        blocks.append('<p>'+html.escape(text)+'</p>')
    def table(headers,rows):
        t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
        for c,h in zip(t.rows[0].cells,headers):c.text=str(h)
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for c,v in zip(t.add_row().cells,row):c.text=str(v)
        blocks.append('<table><thead><tr>'+''.join('<th>'+html.escape(str(h))+'</th>' for h in headers)+
          '</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+html.escape(str(v))+'</td>' for v in row)+'</tr>' for row in rows)+'</tbody></table>')
    def fig(name,caption):
        rel='06_cases/'+name+'.jpg'
        doc.add_picture(str(OUT/rel),width=Cm(17.7))
        doc.paragraphs[-1].paragraph_format.keep_with_next=True
        doc.add_paragraph(caption,style='Caption')
        blocks.append(f'<figure><a href="{rel}"><img src="{rel}" alt="{html.escape(caption)}"></a><figcaption>{html.escape(caption)}</figcaption></figure>')

    doc.add_heading('YOLOv8n-Seg项目场景微调\n及芯片部署模型转换',0)
    blocks.append('<h1>YOLOv8n-Seg项目场景微调及芯片部署模型转换</h1>')
    p('STORY #25028｜父任务 #23579：割草机器人-RGB草坪及作业区域分割开发')
    p('交付日期：2026-10-14；当前冻结模型：B08_v3_2；目标平台：AM970；类别：grass（单类别草坪分割）。')
    p('交付性质：阶段交付。六项主要工作中，第1、3、4、5、6项已完成；第2项“使用独立外场数据完成测试”待补充。本报告不将阶段完成等同于全部验收通过。')
    table(['关键结果','本次实测'],[
        ['PT与ONNX最终二值Mask','50/50一致，平均IoU 100%'],
        ['PT与NPU输出一致性','平均IoU 99.31%；最低96.46%（022）'],
        ['NPU与现有标注','50图平均草坪IoU 98.33%；不是独立外场准确率'],
        ['板端静态推理日志','平均36.01ms，不含完整Mask后处理，不等于摄像头端到端耗时'],
        ['正式模型状态','摄像头版和默认静态版已替换为B08_v3_2；旧模型已备份'],
        ['待补充','独立外场数据、场景评估、连续运行与电源稳定性实测结论']])
    p('阅读索引：第1节任务对照；第2节模型和转换；第3节验证口径；第4节精度改善；第5节典型图片；第6节部署与性能；第7节外场测试补充页；第8节交付及验收说明。')

    heading('1. 六项工作逐项对照',page=True)
    table(['序号','主要工作','状态与证据'],[
        ['1','使用项目增量数据进行模型微调','已完成。训练与数据迭代记录见03_training；最终权重B08_v3_2。'],
        ['2','使用独立外场数据完成测试','待补充。专用补充页见第7节，记录模板见07_field_test_PENDING。'],
        ['3','导出ONNX模型','已完成。固定输入、opset12，50图与PT的最终Mask一致。'],
        ['4','检查工具链对分割算子的支持','已完成当前模型范围内检查。导入、校准及NB导出成功，并完成真实板端推理；不表示任意模型算子均支持。'],
        ['5','转换为芯片NPU模型格式','已完成。UINT8/FP16混合精度NB，模型及配置见02_models。'],
        ['6','对比PyTorch、ONNX和NPU输出','已完成50图三端一致性测试及误差分析，逐图证据见05_evaluation。']])
    p('任务目标中的“项目相机图像分割”已有内部静态验证及摄像头链路验证；独立外场效果仍由第2项补齐。转换前后允许误差尚需项目方确认，本报告不代填验收阈值。')

    heading('2. 模型、配置与转换链路',page=True)
    p('项目历史版本经过标注清理、增量数据和训练策略迭代。B08_v3_2是以yolov8n-seg.pt初始化的项目数据训练版本，不是从B08_v3权重继续微调；项目当前冻结选用该版本，不代表它在每张图或每个指标上都优于所有历史版本。')
    table(['配置','值'],[['输入/训练尺寸','imgsz=640，rect=false，单类别grass'],
      ['训练','AdamW，lr0=0.001，batch=16，epochs=200，patience=20，seed=42'],
      ['数据增强','mosaic=0.2，close_mosaic=100；其他配置见args.yaml'],
      ['日志选点','实际记录139轮，综合Box/Mask最优第119轮；Mask mAP50-95=0.95352'],
      ['ONNX','1×3×640×640；opset12；无内置NMS；固定尺寸；不做简化'],
      ['输出','检测1×37×8400（4坐标+1类别+32系数）；原型1×32×160×160'],
      ['工具链','Acuity 6.33.31；VIP9000NANODI_PLUS_PID0X1000000B'],
      ['模型精度','主体UINT8；坐标链7节点FP16；保留输出拼接FP16；API输出FP32']])
    p('转换流程：PT → ONNX → Acuity导入 → 同一100张校准图重新校准 → 应用混合精度策略 → 导出NB → 板端50图真实输出 → 同一主机后处理对比。每个新模型都重新计算量化参数，未将旧模型的量化数值直接套到新权重。')
    p('前处理配置：add_preproc_node=true，scale=[1/255,1/255,1/255]，reverse_channel=true；后处理配置：add_postproc_node=true，force_float32=true。校准沿用既有inputmeta的fitting=scale及原始100图，不将它描述为与测试图完全相同的Letterbox输入。配置与校准快照均随包保存。')
    p('当前图结构导入后218个节点描述与已验证结构一致。导入/量化/导出日志完整保存；当前转换未出现阻断部署的不可替代算子。PT参考运行时为项目Ultralytics 8.4.128，ONNX导出使用既有转换环境8.3.0；版本信息与输出一致性已记录。')

    heading('3. 验证口径与定量结果',page=True)
    p('本次使用固定内部50图：训练集25图、验证集25图，已用于模型选择与问题分析，不是独立测试集。使用现有多边形标注栅格化、合并grass实例计算像素指标；数据编号和原文件名见sample_index及manifest。不同历史数据/标注版本不能混用比较。')
    table(['口径','设置与用途'],[
      ['A：原生PT模型版本对照','固定640×640 Letterbox；conf=0.25；NMS IoU=0.7；retina_masks=True。用于比较训练版本。'],
      ['B：PT/ONNX/NPU转换一致性','三端共用主机解码；conf≥0.25；NMS IoU>0.45时抑制；Mask>0.5；修正后的floor/ceil裁剪。'],
      ['测试预处理','原图宽640高480，比例1，上下各补80，填充值114；PT/ONNX RGB float32/255；NPU BGR UINT8经NB预处理。'],
      ['指标解释','一致性IoU衡量两端Mask接近程度；标注IoU衡量与现有标注吻合度；均不等同于实例mAP。']])
    p('A口径：B08_v3的50图平均IoU为98.05%，B08_v3_2为98.38%；40张提高、10张下降。B口径的后处理不同，不能直接把A、B两套绝对数值当成同一个实验比较。')
    table(['B口径：三端一致性','均值','最低','检测数一致'],[
      ['PT vs ONNX','100.00%','100.00%','50/50'],['PT vs NPU','99.31%','96.46%（022）','48/50']])
    table(['B口径：NPU对现有标注','B08_v3','B08_v3_2'],[
      ['全部50图平均IoU','98.23%','98.33%'],['其中验证25图平均IoU','98.02%','98.28%'],
      ['018','97.56%','94.92%'],['042','95.55%','96.80%'],['044','95.64%','96.75%'],
      ['总误检像素FP','76,596','62,467'],['总漏检像素FN','40,434','45,850']])
    p('板端新旧版本逐图对照：23张改善、27张下降。总体误检减少，但漏检增加，不能表述为“无精度损失”或“全部图片正常”。当前PT与标注平均IoU为98.49%，NPU为98.33%，转换后仍存在残余损失。')

    heading('4. 精度损失定位与改善方法',page=True)
    heading('4.1 后处理边界裁剪修正',2)
    p('原实现限制框坐标到width-1/height-1并分别对x、y、w、h截断，易导致右/下边界少覆盖像素。修正方法是采用右/下排他边界width/height，左上floor、右下ceil，再裁到合法图像范围。这解决取整裁剪问题，但不能恢复量化造成的真实框坐标或置信度变化。')
    p('009原始框下边界475.75，而PT约478.481；单改裁剪只能补回有限像素。50图一致性均值在三端统一修正后为97.81%，原口径为97.85%，不可声称单独修正裁剪就提高了总体一致性。当前报告中的主机后处理已采用修正；板端NB替换没有同步替换摄像头可执行文件，生产C++逐位一致性仍不是本试验验证范围。')
    heading('4.2 从单节点FP16到坐标计算链FP16',2)
    p('量化后的坐标刻度误差会改变检测框，Mask再被框裁剪时就表现为边缘缺口。先尝试只将最终Mul_2输出升为FP16，效果有限；随后保护Mul_2及必要上游共7个坐标节点。保持其他精度策略、100张校准图和后处理一致，重新校准、导出并上板验证。')
    table(['历史B07_v4隔离实验','平均一致性IoU','最低','日志推理均值'],[
      ['原混合精度＋统一修正裁剪','97.81%','87.29%','约36.05ms'],
      ['仅新增Mul_2 FP16','98.03%','87.72%','35.99ms'],
      ['坐标链7节点FP16','98.89%','89.50%','36.01ms']])
    p('坐标链版41图改善、9图下降。两组实验中类别分数、Mask系数和原型在50图上与基线逐值一致，因此该改善证据针对坐标路径；不能推断分类误差也已消除。约0.06ms的批次差异不应解释为确定加速。')
    for node in ['Mul_2_output_0_6','Concat_5_output_0_11','Div_1_output_0_18','Sub_1_output_0_19','Add_2_output_0_24','Add_1_output_0_26','Sub_output_0_27']:
        p('@model.22/'+node+':out0 → FP16')
    heading('4.3 分类置信度跨阈值与训练版本改善',2)
    p('B08_v3的042右侧实例：PT置信度0.2933，NPU降至0.2242，低于0.25阈值而被过滤；主框止于x=615，裁剪后右侧约25像素宽区域缺失。裁剪前NPU原型仍有草坪响应，不是整幅Mask简单平移，也不是Letterbox错位。')
    p('B08_v3_2保留相同量化策略和0.25阈值，042右侧实例PT置信度0.5567、NPU为0.6523，最终框右边界638.625。右侧x=615:640的预测草坪像素从0恢复为2647。此次改善来自新权重在相同部署策略下的结果，不是新增分类分支FP16，也不是调低阈值；分类量化风险在其他样本上仍存在。')
    p('018道路误检在旧PT本身已出现，因此不能全部归因于NPU量化。通过训练版本迭代，新PT在固定补边输入下显著改善。现有实验不足以单独证明Mosaic是根因；不将训练、标注、随机种子及量化等多个因素混作单因素结论。')

    heading('5. 典型样本改善前后图片',page=True)
    p('以下均使用已有实测输出图，未人工涂改Mask。点击HTML图片或打开06_cases原图可查看大图。红色覆盖图中红色表示预测grass；对标注图中绿色=TP、红色=FP、蓝色=FN。')
    heading('5.1 009：底部被框裁掉的坐标误差',2)
    fig('009_coordinate_chain','图1：左PT、中原混合精度、右7节点FP16。三者都使用修正裁剪。009一致性IoU由96.40%提高到97.78%，框下边界由475.75改善至479.125（PT约478.481）。')
    heading('5.2 018：道路误检的训练版本改善',2)
    fig('018_training_before_after','图2：左至右为标注、B07_v4、B08_v3、B08_v3_2；原生PT固定640×640。018标注IoU由47.52%改善至95.57%。绿色正确、红色误检、蓝色漏检。')
    fig('018_final_backends','图3：最终B08_v3_2原图、PT、ONNX、NPU。大片前景水泥路误检未重现，但仍有局部误检；NPU对标注94.92%，低于前版NPU的97.56%，如实保留该回退。')

    heading('5.3 042：右侧缺失的改善',page=True)
    fig('042_npu_before_after','图4：左标注、中B08_v3 NPU、右B08_v3_2 NPU。旧版右侧蓝色漏检带明显，新版恢复覆盖；标注IoU由95.55%升至96.80%。')
    fig('042_final_backends','图5：最终原图、PT、ONNX、NPU。右侧补充框保留，PT/NPU本图Mask一致性由96.97%升至98.97%，仍有局部边界误差。')
    heading('5.4 044：左下方草坪覆盖改善',2)
    fig('044_training_before_after','图6：标注、B07_v4、B08_v3、B08_v3_2；原生PT同口径。左下方大块漏检得到改善，IoU由87.47%升至97.19%。')
    fig('044_npu_before_after','图7：标注、B08_v3 NPU、B08_v3_2 NPU。当前NPU对标注IoU由95.64%升至96.75%，未宣称所有边缘完全重合。')

    heading('5.5 保留的失败/差异样本',page=True)
    fig('022_remaining_difference','图8：022原图、PT、ONNX、NPU。NPU比PT少一个实例，背景区域覆盖有差异；这是当前最低一致性图（96.46%）。应结合标注核查，不将所有与PT不同的区域直接认定为真实漏检。')
    fig('023_remaining_difference','图9：023原图、PT、ONNX、NPU。仍有实例数量差异，列入外场及后续精度复核清单。')
    p('后续若继续优化，应分离验证分类分支精度、置信度阈值和数据因素，并在完整回归集上同时统计新增误检/漏检。不直接取消框裁剪，也不为修好一张图而未经回归修改正式阈值。')

    heading('6. 板端部署与性能',page=True)
    p('已将摄像头版/root/deploy_mower_cam_vknn/model/network_binary.nb和默认静态版/root/mower_uint8_trial/network_binary.nb替换为本次B08_v3_2，校验值与交付模型一致。旧模型备份：/root/model_backups/before_b08_v3_2_20261014_8421。测试目录为/root/mower_b08_v3_2_20261014。')
    p('模型NB SHA256：'+expected['npu'])
    p('摄像头启动：cd /root/deploy_mower_cam_vknn，然后 sh run_camera.sh left 200。200为帧数，不是秒数；0表示持续运行。图片约每5秒保存一次，结果位于results/<运行目录>/，包含run.log与timing.csv。')
    p('静态50图：日志inference均值36.0124ms，范围35.8854～36.1463ms。静态wrapper计时包含输入提交、推理和输出获取，不含完整Mask后处理；不是NPU内部硬件计时，不等同于整张图片总耗时。')
    table(['室内摄像头链路实测（无检测目标）','代表性记录'],[
      ['处理结果','200帧完成，20.083秒；保存4张彩色图'],
      ['相机/模型初始化','约200ms / 46ms，仅启动阶段'],
      ['采集、模型调用、输出复制','约14ms、36.5ms、40.7ms'],
      ['整轮平均/FPS','约98ms / 10.2FPS'],
      ['限制','画面为室内灯光与展板，无草坪；Mask计算未执行，不能代表有目标或外场FPS']])
    p('摄像头日志model_run_ms单独计时vknn_inference()，与静态wrapper日志的计时边界不同。output_copy_ms约41ms是当前显著开销，不能把整条链路慢归因于NMS。板端曾出现1970年文件时间戳，外场前应校时并记录采集序列；移动电源测试尚无可验收结论。')

    heading('7. 第2项：独立外场测试补充页',page=True)
    p('状态：待补充。该节故意保留，不以内部50图或室内空场景替代。补充时保持已冻结模型和阈值不变；如更换模型，另建版本与报告。')
    table(['待填写项目','记录位置/填写内容'],[
      ['测试日期、地点、人员','_______________________________'],
      ['数据独立性','确认未参与训练、校准及选型；按地点/录像序列与既有数据隔离。结论：________'],
      ['场景覆盖与样本数','晴/阴/逆光/阴影、草路交界、不同草况、裸土/水泥路；记录各组数量：________'],
      ['标注与复核','原图、人工标注、标注规范/复核人：________'],
      ['精度结果','分组IoU、Precision、Recall、FP/FN面积和典型误检/漏检：________'],
      ['连续运行','起止时间、帧数、有效FPS、阶段耗时、异常/重启次数：________'],
      ['移动电源','输出规格、线缆、运行时长、供电中断/重启记录：________'],
      ['验收标准','允许平均/最差单图误差、道路误检约束、性能要求，由项目方确认：________'],
      ['测试结论/签字','通过 / 有条件通过 / 未通过（待实测）；复核人及日期：________']])
    p('补充材料放入07_field_test_PENDING，填写field_records_template.csv，保留相机原图、输出图、run.log、timing.csv。完成后在此节增加分场景结果表及至少3组成功/失败案例，再更新本报告状态。')
    pending=Document()
    pending.add_heading('STORY #25028 第2项独立外场测试补充表',0)
    pending.add_paragraph('当前状态：待补充；不可将内部50图或室内无草坪记录作为独立外场验收。')
    for text in ['测试日期/地点/人员','模型版本与SHA256','数据独立性与序列划分','场景覆盖、原图/标注清单',
                 '各场景IoU/Precision/Recall及误检漏检','典型成功/失败图片及说明','持续FPS与分阶段耗时',
                 '移动电源及异常/重启记录','允许误差与验收门槛','结论、复核人和签字日期']:
        pending.add_heading(text,2);pending.add_paragraph('【待填写】\n\n')
    pending.save(field/'field_test_supplement.docx')

    heading('8. 交付清单与验收边界',page=True)
    table(['目录','内容'],[
      ['01_report.docx / 01_report.html','可编辑Word报告 / 可点击大图的HTML报告'],
      ['02_models','最终PT、ONNX、NB、量化及前后处理配置、校验溯源'],
      ['03_training / 04_conversion','训练配置与日志快照 / 导入量化导出日志'],
      ['05_evaluation','50图逐图指标、三端结果、标注对照、历史实验汇总'],
      ['06_cases','009/018/042/044改善前后及022/023残余差异原尺寸图'],
      ['07_field_test_PENDING','第2项待补充Word表和CSV模板'],
      ['08_deployment','安装脚本、静态测试脚本、室内摄像头日志'],
      ['09_scripts','关键转换、对比、报告脚本快照'],
      ['10_data_snapshot','本次内部50图、GT Mask/匹配标签、校准100图及清单'],
      ['SHA256SUMS.txt','包内文件逐一校验值；不含该清单自身']])
    p('模型、测试输入、校准数据和结果均有SHA256记录。完整训练数据未重复打包，原训练配置包含原环境路径；需要完整重训复现时仍须交接对应版本的完整训练集与工具链。09_scripts保留原工程路径，迁移环境时需调整路径，不是开箱即用的跨平台安装包。')
    p('最终结论：已具备当前冻结模型的训练成果、ONNX、NPU部署模型及内部一致性证据；第1、3、4、5、6项工作完成，第2项待外场补齐。Mask允许误差与安全相关场景门槛由项目方确认，未据平均一致性高而代替外场或实车验收。')
    doc.save(OUT/'01_report.docx')
    page='''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>STORY25028 B08_v3_2阶段交付报告</title><style>
body{font-family:"Microsoft YaHei",sans-serif;color:#20252b;background:white;max-width:1150px;margin:32px auto;padding:0 24px;line-height:1.65}h1,h2{color:#174b46;line-height:1.35}h1{margin-top:42px}table{border-collapse:collapse;width:100%;margin:18px 0;font-size:14px}th,td{border:1px solid #c9d3d1;padding:9px;text-align:left;overflow-wrap:anywhere}th{background:#edf4f2}p{overflow-wrap:anywhere}figure{margin:20px 0}img{max-width:100%;height:auto}figcaption{font-size:14px;color:#48515a}a{color:#145c82}@media print{h1{break-before:page}figure, tr{break-inside:avoid}body{margin:0;max-width:none}}
</style><body>'''+''.join(blocks)+'</body></html>'
    (OUT/'01_report.html').write_text(page,encoding='utf-8')
    (OUT/'START_HERE.txt').write_text('STORY #25028 阶段交付包，2026-10-14\n'
       '先打开01_report.docx（可编辑），或01_report.html（点击图片查看大图）。\n'
       '第1、3、4、5、6项完成；第2项外场测试待补充，模板在07_field_test_PENDING。\n'
       '当前交付模型B08_v3_2，包含真实改善和残余误差，不宣称全部验收通过。\n',encoding='utf-8')
    files=sorted(p for p in OUT.rglob('*') if p.is_file() and p.name!='SHA256SUMS.txt')
    (OUT/'SHA256SUMS.txt').write_text(''.join(f'{sha(p)}  {p.relative_to(OUT).as_posix()}\n' for p in files))
    archive=OUT.with_suffix('.zip')
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED,compresslevel=6,strict_timestamps=False) as z:
        for p in sorted(OUT.rglob('*')):
            if p.is_file():z.write(p,str(Path(OUT.name)/p.relative_to(OUT)))
    print(json.dumps(dict(folder=str(OUT),archive=str(archive),file_count=len(files)+1,
                         zip_bytes=archive.stat().st_size,matching_labels=sum(s['label_snapshot_matches'] for s in snapshots)),indent=2))


if __name__=='__main__':
    main()
